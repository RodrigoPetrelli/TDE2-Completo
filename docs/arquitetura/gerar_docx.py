"""Preenche o template oficial arc42 (.docx, PT) com o conteúdo de gerado/arc42.html.

Uso: python gerar_docx.py <arc42.html preenchido> <template.docx> <grupo.json> <saida.docx>

Mantém os estilos, a capa e a seção "Sobre o arc42" do template e substitui as
seções-modelo pelo conteúdo do PedeJá, usando os estilos do próprio template
(Heading 1-3, First Paragraph, Body Text, Compact, Definition Term/Definition,
Image Caption, Source Code, Table).
"""

import json
import re
import struct
import sys
from datetime import date
from pathlib import Path

from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1D, 0x4F, 0x91)
MARGEM = Cm(2.0)
MARGEM_PAISAGEM = Cm(1.4)
A4 = (Cm(21.0), Cm(29.7))
ALTURA_MAX = {False: Cm(20.5), True: Cm(14.2)}  # imagem: retrato / paisagem
MESES = ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro', 'outubro',
         'novembro', 'dezembro']
NUM_ID_BULLET = 90


class Gerador:
    def __init__(self, html: Path, template: Path, grupo: dict):
        self.dir_imagens = html.parent
        self.soup = BeautifulSoup(html.read_text(encoding='utf-8'), 'html.parser')
        self.doc = Document(str(template))
        self.grupo = grupo
        self.paisagem = False
        self.apos_titulo = False
        self.inicio_de_secao = False
        self.retorno_retrato_pendente = False
        self.forcar_quebra = False

    # --- template -----------------------------------------------------------
    def preparar_template(self):
        corpo = self.doc.element.body
        paragrafos = self.doc.paragraphs
        inicio = next(p for p in paragrafos if p.style.name == 'Heading 1' and p.text.strip() == 'Introdução e Objetivos')
        remover = False
        for elemento in list(corpo):
            if elemento is inicio._p:
                remover = True
            if remover and elemento.tag != qn('w:sectPr'):
                corpo.remove(elemento)

        secao = self.doc.sections[0]
        secao.page_width, secao.page_height = A4
        secao.orientation = WD_ORIENT.PORTRAIT
        for lado in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
            setattr(secao, lado, MARGEM)
        secao.different_first_page_header_footer = True
        self.rodape(secao)

        self.capa()
        self.numeracao_bullets()

    def capa(self):
        g = self.grupo
        titulo = next(p for p in self.doc.paragraphs if p.style.name == 'Title')
        data = next(p for p in self.doc.paragraphs if p.style.name == 'Date')
        titulo.text = 'PedeJá'
        titulo.paragraph_format.space_before = Cm(6)
        titulo.runs[0].font.size = Pt(40)
        titulo.runs[0].bold = True
        titulo.runs[0].font.color.rgb = AZUL
        hoje = date.today()
        data.text = f'Versão 1.0 · {hoje.day} de {MESES[hoje.month - 1]} de {hoje.year}'

        def antes_da_data(texto, estilo, negrito=False):
            p = data.insert_paragraph_before(style=estilo)
            p.add_run(texto).bold = negrito
            return p

        antes_da_data('Documentação de Arquitetura de Software (arc42)', 'Subtitle')
        antes_da_data(f"{g['disciplina']} — {g['atividade']}", 'Author')
        if g.get('instituicao'):
            antes_da_data(g['instituicao'], 'Author')
        if g.get('professor'):
            antes_da_data(f"Professor(a): {g['professor']}", 'Author')
        antes_da_data('Integrantes do grupo', 'Author', negrito=True).paragraph_format.space_before = Cm(1.5)
        for nome in g['integrantes']:
            antes_da_data(nome, 'Author').runs[0].font.size = Pt(13)
        antes_da_data('arc42 · C4 Model (níveis 1 a 4) · UML · DER · Software Architecture Canvas',
                      'Author').paragraph_format.space_before = Cm(1.5)

        # O título vazio do template vira apenas a quebra de página antes de "Sobre o arc42"
        vazio = next(p for p in self.doc.paragraphs if p.style.name == 'Heading 1' and not p.text.strip())
        vazio.style = self.doc.styles['Normal']
        vazio.paragraph_format.page_break_before = True

        # Sumário (atualizado pelo Word na exportação)
        p = self.doc.add_paragraph('Sumário', style='TOC Heading')
        p.paragraph_format.page_break_before = True
        self.campo(self.doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', 'Atualize o sumário (F9).')

    def rodape(self, secao):
        p = secao.footer.paragraphs[0]
        p.text = ''
        p.paragraph_format.tab_stops.add_tab_stop(A4[0] - 2 * MARGEM, WD_TAB_ALIGNMENT.RIGHT)
        p.add_run('PedeJá — Documentação de Arquitetura (arc42)').font.size = Pt(8)
        p.add_run('\t').font.size = Pt(8)
        r = p.add_run('Página ')
        r.font.size = Pt(8)
        self.campo(p, 'PAGE', '1', tamanho=8)
        p.add_run(' de ').font.size = Pt(8)
        self.campo(p, 'NUMPAGES', '1', tamanho=8)

    @staticmethod
    def campo(paragrafo, instrucao, texto, tamanho=None):
        def run_com(filho):
            r = OxmlElement('w:r')
            if tamanho:
                rpr = OxmlElement('w:rPr')
                sz = OxmlElement('w:sz')
                sz.set(qn('w:val'), str(tamanho * 2))
                rpr.append(sz)
                r.append(rpr)
            r.append(filho)
            paragrafo._p.append(r)

        inicio = OxmlElement('w:fldChar')
        inicio.set(qn('w:fldCharType'), 'begin')
        run_com(inicio)
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = f' {instrucao} '
        run_com(instr)
        sep = OxmlElement('w:fldChar')
        sep.set(qn('w:fldCharType'), 'separate')
        run_com(sep)
        t = OxmlElement('w:t')
        t.text = texto
        run_com(t)
        fim = OxmlElement('w:fldChar')
        fim.set(qn('w:fldCharType'), 'end')
        run_com(fim)

    def numeracao_bullets(self):
        numeracao = self.doc.part.numbering_part.element
        abstrato = OxmlElement('w:abstractNum')
        abstrato.set(qn('w:abstractNumId'), str(NUM_ID_BULLET))
        for nivel in range(2):
            lvl = OxmlElement('w:lvl')
            lvl.set(qn('w:ilvl'), str(nivel))
            for tag, valor in (('w:start', '1'), ('w:numFmt', 'bullet'), ('w:lvlText', '•' if nivel == 0 else '◦'),
                               ('w:lvlJc', 'left')):
                e = OxmlElement(tag)
                e.set(qn('w:val'), valor)
                lvl.append(e)
            ppr = OxmlElement('w:pPr')
            ind = OxmlElement('w:ind')
            ind.set(qn('w:left'), str(360 * (nivel + 1)))
            ind.set(qn('w:hanging'), '270')
            ppr.append(ind)
            lvl.append(ppr)
            abstrato.append(lvl)
        primeiro_num = numeracao.find(qn('w:num'))
        if primeiro_num is not None:
            primeiro_num.addprevious(abstrato)
        else:
            numeracao.append(abstrato)
        num = OxmlElement('w:num')
        num.set(qn('w:numId'), str(NUM_ID_BULLET))
        ref = OxmlElement('w:abstractNumId')
        ref.set(qn('w:val'), str(NUM_ID_BULLET))
        num.append(ref)
        numeracao.append(num)

    # --- conteúdo ---------------------------------------------------------
    def gerar(self):
        self.preparar_template()
        for no in self.soup.body.children:
            self.bloco(no, self.doc)

    def bloco(self, no, destino):
        if isinstance(no, Comment) or not isinstance(no, Tag):
            return
        classes = no.get('class') or []
        if no.name == 'section' and ({'capa', 'sumario'} & set(classes)):
            return
        paisagem = no.name == 'div' and 'paisagem' in classes
        if self.retorno_retrato_pendente and not paisagem:
            # Volta ao retrato só quando houver conteúdo, evitando páginas em branco
            # entre dois trechos em paisagem consecutivos.
            self.retorno_retrato_pendente = False
            self.trocar_orientacao(False)
        if no.name in ('h2', 'h3', 'h4'):
            nivel = {'h2': 1, 'h3': 2, 'h4': 3}[no.name]
            p = destino.add_paragraph(self.texto(no), style=f'Heading {nivel}')
            if (nivel == 1 and not self.inicio_de_secao) or self.forcar_quebra or 'nova-pagina' in classes:
                p.paragraph_format.page_break_before = True
            self.inicio_de_secao = False
            self.forcar_quebra = False
            self.apos_titulo = True
            return
        self.inicio_de_secao = False
        if no.name == 'p':
            p = destino.add_paragraph(style='First Paragraph' if self.apos_titulo else 'Body Text')
            self.inline(p, no)
        elif no.name == 'ul':
            for li in no.find_all('li', recursive=False):
                self.item_lista(destino, li)
        elif no.name == 'table':
            self.tabela(no, destino)
        elif no.name == 'figure':
            self.figura(no, destino)
        elif no.name == 'dl':
            for filho in no.find_all(['dt', 'dd'], recursive=False):
                p = destino.add_paragraph(style='Definition Term' if filho.name == 'dt' else 'Definition')
                self.inline(p, filho)
        elif no.name == 'pre':
            self.arvore(no, destino)
        elif paisagem:
            if self.retorno_retrato_pendente:
                # Paisagem seguida de paisagem: continua na mesma seção, em nova página
                self.retorno_retrato_pendente = False
                self.forcar_quebra = True
            else:
                self.trocar_orientacao(True)
            for filho in no.children:
                self.bloco(filho, destino)
            self.retorno_retrato_pendente = True
            return
        elif no.name == 'div' and 'canvas' in classes:
            self.canvas(no, destino)
        else:
            for filho in no.children:
                self.bloco(filho, destino)
            return
        self.apos_titulo = False

    def trocar_orientacao(self, paisagem):
        secao = self.doc.add_section(WD_SECTION.NEW_PAGE)
        largura, altura = A4
        secao.orientation = WD_ORIENT.LANDSCAPE if paisagem else WD_ORIENT.PORTRAIT
        secao.page_width, secao.page_height = (altura, largura) if paisagem else (largura, altura)
        for lado in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
            setattr(secao, lado, MARGEM_PAISAGEM if paisagem else MARGEM)
        # Só a capa (primeira seção) dispensa rodapé; as demais seções herdariam essa opção.
        secao.different_first_page_header_footer = False
        self.paisagem = paisagem
        self.inicio_de_secao = True

    @staticmethod
    def texto(no):
        return re.sub(r'\s+', ' ', no.get_text()).strip()

    def inline(self, paragrafo, no, tamanho=None):
        segmentos = []

        def coletar(n, negrito=False, italico=False, codigo=False):
            if isinstance(n, Comment):
                return
            if isinstance(n, NavigableString):
                segmentos.append([re.sub(r'\s+', ' ', str(n)), negrito, italico, codigo])
                return
            if n.name == 'br':
                segmentos.append(['\n', negrito, italico, codigo])
                return
            for filho in n.children:
                coletar(filho, negrito or n.name in ('strong', 'b'), italico or n.name in ('em', 'i'),
                        codigo or n.name == 'code')

        coletar(no)
        segmentos = [s for s in segmentos if s[0]]
        if not segmentos:
            return
        segmentos[0][0] = segmentos[0][0].lstrip()
        segmentos[-1][0] = segmentos[-1][0].rstrip()
        for texto, negrito, italico, codigo in segmentos:
            if not texto:
                continue
            run = paragrafo.add_run(texto)
            run.bold = negrito or None
            run.italic = italico or None
            if codigo:
                run.font.name = 'Consolas'
                run.font.size = Pt((tamanho or 10) - 1)
                run.font.color.rgb = RGBColor(0x5A, 0x1E, 0x8C)
            elif tamanho:
                run.font.size = Pt(tamanho)

    def item_lista(self, destino, li, tamanho=None):
        p = destino.add_paragraph(style='Compact')
        ppr = p._p.get_or_add_pPr()
        numpr = OxmlElement('w:numPr')
        ilvl = OxmlElement('w:ilvl')
        ilvl.set(qn('w:val'), '0')
        numid = OxmlElement('w:numId')
        numid.set(qn('w:val'), str(NUM_ID_BULLET))
        numpr.append(ilvl)
        numpr.append(numid)
        ppr.append(numpr)
        self.inline(p, li, tamanho)
        return p

    # --- tabelas ----------------------------------------------------------
    @staticmethod
    def bordas(tabela, cor='C9D3E0', tamanho=4):
        tblpr = tabela._tbl.tblPr
        bordas = OxmlElement('w:tblBorders')
        for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            b = OxmlElement(f'w:{lado}')
            b.set(qn('w:val'), 'single' if tamanho else 'nil')
            b.set(qn('w:sz'), str(tamanho))
            b.set(qn('w:color'), cor)
            bordas.append(b)
        tblpr.append(bordas)
        largura = OxmlElement('w:tblW')
        largura.set(qn('w:type'), 'pct')
        largura.set(qn('w:w'), '5000')
        existente = tblpr.find(qn('w:tblW'))
        if existente is not None:
            tblpr.remove(existente)
        tblpr.append(largura)

    @staticmethod
    def sombrear(celula, cor):
        tcpr = celula._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), cor)
        tcpr.append(shd)

    def tabela(self, no, destino):
        linhas = no.find_all('tr')
        colunas = max(sum(int(c.get('colspan', 1)) for c in tr.find_all(['th', 'td'], recursive=False)) for tr in linhas)
        tabela = destino.add_table(rows=len(linhas), cols=colunas)
        tabela.style = self.doc.styles['Table']
        self.bordas(tabela)
        for i, tr in enumerate(linhas):
            cabecalho = tr.parent.name == 'thead'
            if cabecalho:
                trpr = tabela.rows[i]._tr.get_or_add_trPr()
                trpr.append(OxmlElement('w:tblHeader'))
            j = 0
            for c in tr.find_all(['th', 'td'], recursive=False):
                span = int(c.get('colspan', 1))
                celula = tabela.cell(i, j)
                if span > 1:
                    celula = celula.merge(tabela.cell(i, j + span - 1))
                p = celula.paragraphs[0]
                p.style = self.doc.styles['Compact']
                self.inline(p, c, tamanho=8.5)
                if cabecalho:
                    self.sombrear(celula, '1D4F91')
                    for r in p.runs:
                        r.bold = True
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                elif c.name == 'th':
                    self.sombrear(celula, 'EAF1FA')
                    for r in p.runs:
                        r.bold = True
                j += span
        # Sem larguras fixas: o Word ajusta as colunas ao conteúdo (tabela a 100% da página)
        for tcw in tabela._tbl.iter(qn('w:tcW')):
            tcw.getparent().remove(tcw)
        destino.add_paragraph(style='Body Text')

    def canvas(self, no, destino):
        blocos = no.find_all('div', recursive=False)
        tabela = destino.add_table(rows=(len(blocos) + 2) // 3, cols=3)
        tabela.style = self.doc.styles['Table']
        self.bordas(tabela, cor='1D4F91', tamanho=12)
        for linha in tabela.rows:
            linha._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for k, bloco in enumerate(blocos):
            celula = tabela.cell(k // 3, k % 3)
            self.sombrear(celula, 'FBFCFE')
            primeiro = True
            for filho in bloco.find_all(['h4', 'p', 'ul'], recursive=False):
                if filho.name == 'h4':
                    p = celula.paragraphs[0] if primeiro else celula.add_paragraph()
                    p.style = self.doc.styles['Compact']
                    r = p.add_run(self.texto(filho).upper())
                    r.bold = True
                    r.font.size = Pt(9)
                    r.font.color.rgb = AZUL
                elif filho.name == 'p':
                    p = celula.add_paragraph(style='Compact')
                    self.inline(p, filho, tamanho=8.5)
                else:
                    for li in filho.find_all('li', recursive=False):
                        self.item_lista(celula, li, tamanho=8.5)
                primeiro = False
        destino.add_paragraph(style='Body Text')

    def arvore(self, no, destino):
        linhas = no.get_text().rstrip('\n').split('\n')
        meio = (len(linhas) + 1) // 2
        tabela = destino.add_table(rows=1, cols=2)
        self.bordas(tabela, tamanho=0)
        for coluna, parte in enumerate((linhas[:meio], linhas[meio:])):
            celula = tabela.cell(0, coluna)
            self.sombrear(celula, 'F6F8FA')
            for k, linha in enumerate(parte):
                p = celula.paragraphs[0] if k == 0 else celula.add_paragraph()
                p.style = self.doc.styles['Source Code']
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                r = p.add_run(linha)
                r.font.name = 'Consolas'
                r.font.size = Pt(7)
        destino.add_paragraph(style='Body Text')

    # --- figuras ----------------------------------------------------------
    def figura(self, no, destino):
        img = no.find('img')
        caminho = self.dir_imagens / img['src']
        porcentagem = int(re.search(r'width:(\d+)%', img.get('style', 'width:100%')).group(1)) / 100
        largura_util = A4[1] - 2 * MARGEM_PAISAGEM if self.paisagem else A4[0] - 2 * MARGEM
        with open(caminho, 'rb') as f:
            cabecalho = f.read(24)
        largura_px, altura_px = struct.unpack('>II', cabecalho[16:24])
        largura = int(largura_util * porcentagem)
        altura = int(largura * altura_px / largura_px)
        if altura > ALTURA_MAX[self.paisagem]:
            largura = int(largura * ALTURA_MAX[self.paisagem] / altura)

        p = destino.add_paragraph(style='Captioned Figure')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(caminho), width=largura)
        legenda = destino.add_paragraph(self.texto(no.find('figcaption')), style='Image Caption')
        legenda.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main():
    html, template, grupo, saida = map(Path, sys.argv[1:5])
    gerador = Gerador(html, template, json.loads(grupo.read_text(encoding='utf-8')))
    gerador.gerar()
    gerador.doc.core_properties.title = 'PedeJá — Documentação de Arquitetura (arc42)'
    gerador.doc.core_properties.author = ', '.join(gerador.grupo['integrantes'])
    gerador.doc.save(str(saida))
    print(f'DOCX gerado: {saida}')


if __name__ == '__main__':
    main()
